#!/usr/bin/env python3
"""Build normalized B-group JSONL and Markdown reports after a clean audit."""

from __future__ import annotations

import json
import subprocess
import sys
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
OUTPUT_ROOT = ROOT / "outputs" / "b_group_gt_repair_20260817"
UNBLOCK_ROOT = ROOT / "outputs" / "b_group_eval_unblock_20260818"
WORKBOOK = ROOT / "181条待修复GT_A组_B组_20260817.xlsx"
AUDITOR = OUTPUT_ROOT / "audit_reports.py"
JSONL_OUTPUT = OUTPUT_ROOT / "B组_96条_GT修复最终报告.jsonl"
MARKDOWN_OUTPUT = OUTPUT_ROOT / "B组_96条_GT修复最终报告.md"


def load_expected() -> list[tuple[str, str]]:
    workbook = load_workbook(WORKBOOK, read_only=True, data_only=True)
    return [
        (str(task_id), str(snapshot))
        for task_id, snapshot in workbook["B组"].iter_rows(
            min_row=2, max_col=2, values_only=True
        )
        if task_id and snapshot
    ]


def discover_tasks() -> dict[str, dict]:
    tasks = {}
    for path in ROOT.glob("task/**/task-*.json"):
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, UnicodeDecodeError, json.JSONDecodeError):
            continue
        task_id = data.get("id")
        if isinstance(task_id, str):
            tasks[task_id] = data
    return tasks


def load_records() -> dict[str, dict]:
    records = {}
    for path in sorted(OUTPUT_ROOT.glob("*/report.jsonl")):
        for line in path.read_text(encoding="utf-8").splitlines():
            if not line.strip():
                continue
            record = json.loads(line)
            records[record["task_id"]] = record
    for path in sorted(UNBLOCK_ROOT.glob("*/report.jsonl")):
        for line in path.read_text(encoding="utf-8").splitlines():
            if not line.strip():
                continue
            record = json.loads(line)
            records[record["task_id"]] = record
    return records


def normalize_hashes(record: dict) -> dict[str, dict[str, str | None]]:
    result = {}
    before_after = record.get("sha256_before_after")
    if isinstance(before_after, dict):
        for path, values in before_after.items():
            if isinstance(values, dict):
                result[path] = {"before": values.get("before"), "after": values.get("after")}

    hashes = record.get("hashes")
    if isinstance(hashes, list):
        for values in hashes:
            if isinstance(values, dict) and isinstance(values.get("path"), str):
                result[values["path"]] = {
                    "before": values.get("before_sha256"),
                    "after": values.get("after_sha256"),
                }

    repository_hashes = record.get("repository_hashes")
    if isinstance(repository_hashes, dict):
        for path, values in repository_hashes.items():
            if isinstance(values, dict):
                result[path] = {"before": values.get("before"), "after": values.get("after")}
    elif isinstance(repository_hashes, list):
        for values in repository_hashes:
            if isinstance(values, dict) and isinstance(values.get("path"), str):
                result[values["path"]] = {
                    "before": values.get("before_sha256", values.get("before")),
                    "after": values.get("after_sha256", values.get("after")),
                }

    before = record.get("sha256_before")
    after = record.get("sha256_after")
    if isinstance(before, dict) or isinstance(after, dict):
        before = before if isinstance(before, dict) else {}
        after = after if isinstance(after, dict) else {}
        for path in before.keys() | after.keys():
            result[path] = {"before": before.get(path), "after": after.get(path)}
    return result


def normalize_eval(record: dict, task: dict) -> dict:
    configured_command = task.get("evaluator", {}).get("result", {}).get("command")
    direct = record.get("eval_result")
    if isinstance(direct, dict) and direct.get("returncode") is not None:
        stdout = direct.get("stdout", direct.get("output"))
        # Early cli2 reports used the display literal ``True\\n``. Recover the
        # actual command output from the persisted production response.
        if stdout == r"True\n":
            for raw_path in normalize_logs(record):
                if not raw_path.endswith("eval_response.json"):
                    continue
                try:
                    response = json.loads((ROOT / raw_path).read_text(encoding="utf-8"))
                except (OSError, UnicodeDecodeError, json.JSONDecodeError):
                    continue
                if response.get("returncode") == 0 and isinstance(response.get("output"), str):
                    stdout = response["output"]
                    break
        score = direct.get("score")
        if score is None:
            score = 1.0 if record.get("status") == "passed" else 0.0
        return {
            "command": direct.get("command", configured_command),
            "returncode": direct.get("returncode"),
            "score": score,
            "stdout": stdout,
            "details": direct,
        }
    snapshot = record["snapshot"]
    if snapshot == "ANSYS-ABAQUS-AUTOCAD":
        return {
            "command": record.get("eval_command", configured_command),
            "returncode": record.get("eval_returncode"),
            "score": record.get("score"),
            "stdout": record.get("eval_stdout"),
        }

    raw = record.get("eval") if isinstance(record.get("eval"), dict) else record.get("eval_result", {})
    if not isinstance(raw, dict):
        raw = {}
    production = raw.get("production") if isinstance(raw.get("production"), dict) else raw
    score = raw.get("score", production.get("score"))
    if score is None:
        score = 1.0 if record.get("status") == "passed" else 0.0
    return {
        "command": raw.get("command", configured_command),
        "returncode": production.get("returncode"),
        "score": score,
        "stdout": production.get("stdout", raw.get("output")),
        "details": raw,
    }


def normalize_logs(record: dict) -> list[str]:
    logs = []
    for key in ("logs", "evidence_logs"):
        raw = record.get(key)
        if isinstance(raw, str):
            logs.append(raw)
        elif isinstance(raw, list):
            logs.extend(value for value in raw if isinstance(value, str))
        elif isinstance(raw, dict):
            logs.extend(value for value in raw.values() if isinstance(value, str))
    for key in ("log", "diagnosis_log"):
        value = record.get(key)
        if isinstance(value, str):
            logs.append(value)
    return sorted(set(logs))


def normalize_record(record: dict, task: dict) -> dict:
    logs = normalize_logs(record)
    disconnect = UNBLOCK_ROOT / record["snapshot"] / "disconnect_verification.json"
    if disconnect.is_file():
        logs.append(disconnect.relative_to(ROOT).as_posix())
        logs = sorted(set(logs))
    return {
        "task_id": record["task_id"],
        "snapshot": record["snapshot"],
        "failure_reason": record.get("failure_reason"),
        "instruction": task.get("instruction"),
        "production_staging": record.get(
            "production_staging", record.get("instruction_eval_contract", record.get("instruction_and_staging_review"))
        ),
        "modified_files": record.get("modified_files", []),
        "sha256_before_after": normalize_hashes(record),
        "real_software_generation": record.get("real_software_generation"),
        "eval_result": normalize_eval(record, task),
        "logs": logs,
        "cleanup": record.get("cleanup"),
        "status": record.get("status"),
    }


def compact(value: object) -> str:
    if value is None:
        return "none"
    if isinstance(value, str):
        return value.replace("\n", " ").strip()
    return json.dumps(value, ensure_ascii=False, sort_keys=True)


def render_markdown(records: list[dict]) -> str:
    status_counts = Counter(record["status"] for record in records)
    by_snapshot = defaultdict(Counter)
    for record in records:
        by_snapshot[record["snapshot"]][record["status"]] += 1
    modified_records = [record for record in records if record["modified_files"]]
    modified_file_count = sum(len(record["modified_files"]) for record in modified_records)

    lines = [
        "# B组 96 条 GT 修复最终报告",
        "",
        "## 汇总",
        "",
        f"- 完成数：{len(records)}/96",
        f"- 通过数：{status_counts['passed']}",
        f"- Eval 阻塞数：{status_counts['eval_blocked']}",
        f"- 环境阻塞数：{status_counts['environment_blocked']}",
        f"- 修改 Task 数：{len(modified_records)}",
        f"- 修改文件数：{modified_file_count}",
        "",
        "| Snapshot | 完成 | passed | eval_blocked | environment_blocked |",
        "|---|---:|---:|---:|---:|",
    ]
    for snapshot in sorted(by_snapshot):
        counts = by_snapshot[snapshot]
        lines.append(
            f"| {snapshot} | {sum(counts.values())} | {counts['passed']} | "
            f"{counts['eval_blocked']} | {counts['environment_blocked']} |"
        )

    lines.extend(["", "## 逐 Task 记录", ""])
    for index, record in enumerate(records, 1):
        eval_result = record["eval_result"]
        lines.extend(
            [
                f"### {index}. {record['task_id']}",
                "",
                f"- Snapshot：`{record['snapshot']}`",
                f"- 最终状态：`{record['status']}`",
                f"- 失败原因：{compact(record['failure_reason'])}",
                f"- 修改文件：{compact(record['modified_files'])}",
                f"- 真实软件生成过程：{compact(record['real_software_generation'])}",
                f"- Eval 命令：`{compact(eval_result.get('command'))}`",
                f"- Eval 返回值：`{compact(eval_result.get('returncode'))}`",
                f"- Eval 分数：`{compact(eval_result.get('score'))}`",
                f"- 日志位置：{compact(record['logs'])}",
                f"- 清理动作：{compact(record['cleanup'])}",
                "",
                "| 文件 | 修改前 SHA-256 | 修改后 SHA-256 |",
                "|---|---|---|",
            ]
        )
        hashes = record["sha256_before_after"]
        if hashes:
            for path, values in hashes.items():
                lines.append(
                    f"| `{path}` | `{values.get('before') or 'absent'}` | "
                    f"`{values.get('after') or 'absent'}` |"
                )
        else:
            lines.append("| none | `absent` | `absent` |")
        lines.append("")
    return "\n".join(lines)


def main() -> int:
    audit = subprocess.run([sys.executable, str(AUDITOR)], cwd=ROOT, check=False)
    if audit.returncode != 0:
        print("Refusing to build final reports because audit_reports.py failed.", file=sys.stderr)
        return audit.returncode

    expected = load_expected()
    tasks = discover_tasks()
    source_records = load_records()
    normalized = [normalize_record(source_records[task_id], tasks[task_id]) for task_id, _ in expected]

    JSONL_OUTPUT.write_text(
        "".join(json.dumps(record, ensure_ascii=False) + "\n" for record in normalized),
        encoding="utf-8",
    )
    MARKDOWN_OUTPUT.write_text(render_markdown(normalized), encoding="utf-8")
    print(JSONL_OUTPUT.relative_to(ROOT))
    print(MARKDOWN_OUTPUT.relative_to(ROOT))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
