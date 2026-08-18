#!/usr/bin/env python3
"""Read-only consistency audit for the B-group GT repair reports."""

from __future__ import annotations

import hashlib
import json
import re
import subprocess
import sys
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
OUTPUT_ROOT = ROOT / "outputs" / "b_group_gt_repair_20260817"
UNBLOCK_ROOT = ROOT / "outputs" / "b_group_eval_unblock_20260818"
WORKBOOK = ROOT / "181条待修复GT_A组_B组_20260817.xlsx"
ALLOWED_STATUSES = {"passed", "eval_blocked", "environment_blocked"}
SHA256_RE = re.compile(r"^[0-9a-f]{64}$")


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def git_file_sha256_candidates(relative: Path) -> set[str]:
    result = subprocess.run(
        ["git", "show", f"HEAD:{relative.as_posix()}"],
        cwd=ROOT,
        stdout=subprocess.PIPE,
        stderr=subprocess.DEVNULL,
        check=False,
    )
    if result.returncode != 0:
        return set()

    candidates = {hashlib.sha256(result.stdout).hexdigest()}
    # A report may intentionally record the SHA-256 of the materialized LFS
    # object rather than the hash of the small pointer stored in Git.
    match = re.fullmatch(
        rb"version https://git-lfs\.github\.com/spec/v1\r?\n"
        rb"oid sha256:([0-9a-f]{64})\r?\n"
        rb"size [0-9]+\r?\n?",
        result.stdout,
    )
    if match:
        candidates.add(match.group(1).decode("ascii"))
    return candidates


def git_task_changes() -> set[str]:
    result = subprocess.run(
        ["git", "status", "--porcelain=v1", "-z", "--untracked-files=all", "--", "task"],
        cwd=ROOT,
        stdout=subprocess.PIPE,
        check=True,
    )
    paths = set()
    entries = result.stdout.decode("utf-8", errors="surrogateescape").split("\0")
    index = 0
    while index < len(entries):
        entry = entries[index]
        index += 1
        if not entry:
            continue
        status = entry[:2]
        paths.add(entry[3:])
        if "R" in status or "C" in status:
            if index < len(entries) and entries[index]:
                paths.add(entries[index])
                index += 1
    return paths


def load_expected() -> list[tuple[str, str]]:
    workbook = load_workbook(WORKBOOK, read_only=True, data_only=True)
    worksheet = workbook["B组"]
    rows = []
    for task_id, snapshot in worksheet.iter_rows(
        min_row=2, max_col=2, values_only=True
    ):
        if task_id and snapshot:
            rows.append((str(task_id), str(snapshot)))
    return rows


def discover_tasks() -> dict[str, tuple[Path, dict]]:
    tasks = {}
    duplicates = defaultdict(list)
    for path in ROOT.glob("task/**/task-*.json"):
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, UnicodeDecodeError, json.JSONDecodeError):
            continue
        task_id = data.get("id")
        if not isinstance(task_id, str):
            continue
        duplicates[task_id].append(path)
        tasks[task_id] = (path.parent, data)
    duplicate_ids = {key: value for key, value in duplicates.items() if len(value) > 1}
    if duplicate_ids:
        formatted = ", ".join(
            f"{task_id} ({len(paths)})" for task_id, paths in sorted(duplicate_ids.items())
        )
        raise RuntimeError(f"duplicate task definitions: {formatted}")
    return tasks


def load_reports() -> tuple[list[dict], list[str]]:
    records = {}
    errors = []
    for layer_name, root in (("baseline", OUTPUT_ROOT), ("unblock", UNBLOCK_ROOT)):
        seen_in_layer = set()
        for path in sorted(root.glob("*/report.jsonl")):
            for line_number, line in enumerate(path.read_text(encoding="utf-8").splitlines(), 1):
                if not line.strip():
                    continue
                try:
                    record = json.loads(line)
                except json.JSONDecodeError as exc:
                    errors.append(f"{path.relative_to(ROOT)}:{line_number}: invalid JSON: {exc}")
                    continue
                task_id = record.get("task_id")
                if task_id in seen_in_layer:
                    errors.append(f"{task_id}: duplicate {layer_name} report record")
                    continue
                seen_in_layer.add(task_id)
                record["_report_path"] = str(path.relative_to(ROOT))
                record["_line_number"] = line_number
                records[task_id] = record
    return list(records.values()), errors


def path_strings(value: object):
    if isinstance(value, str):
        if value.startswith("outputs/"):
            yield value
    elif isinstance(value, dict):
        for child in value.values():
            yield from path_strings(child)
    elif isinstance(value, list):
        for child in value:
            yield from path_strings(child)


def eval_pass_evidence(record: dict) -> bool:
    direct = record.get("eval_result")
    if isinstance(direct, dict) and direct.get("returncode") == 0:
        stdout = str(direct.get("stdout", direct.get("output", ""))).replace("\r\n", "\n")
        score = direct.get("score")
        if stdout == "True\n" and (
            direct.get("exact_match") is True
            or (isinstance(score, (int, float)) and score > 0)
        ):
            return True
        if str(record.get("task_id", "")).startswith("v-quantified-"):
            try:
                returned_score = float(stdout.strip())
            except (TypeError, ValueError):
                return False
            return (
                isinstance(score, (int, float))
                and score > 0
                and abs(returned_score - float(score)) <= 1.0e-6
            )
    snapshot = record.get("snapshot")
    if snapshot in {"cli2-archicad27-openstudio310-win", "Blender-4.2.3"}:
        result = record.get("eval_result", {})
        stdout = str(result.get("stdout", "")).replace("\r\n", "\n")
        if result.get("returncode") != 0:
            return False
        if stdout == "True\n":
            return True
        if stdout != r"True\n":
            return False
        for raw_path in path_strings(record):
            if not raw_path.endswith("eval_response.json"):
                continue
            path = ROOT / raw_path
            try:
                response = json.loads(path.read_text(encoding="utf-8"))
            except (OSError, UnicodeDecodeError, json.JSONDecodeError):
                continue
            raw_stdout = str(response.get("output", "")).replace("\r\n", "\n")
            if response.get("returncode") == 0 and raw_stdout == "True\n":
                return True
        return False
    if snapshot == "ANSYS-2026R":
        result = record.get("eval", {})
        score = result.get("score")
        return result.get("returncode") == 0 and isinstance(score, (int, float)) and score > 0
    if snapshot == "ANSYS-ABAQUS-AUTOCAD":
        score = record.get("score")
        return record.get("eval_returncode") == 0 and isinstance(score, (int, float)) and score > 0
    if snapshot == "ABAQUS-2025L":
        result = record.get("eval", {})
        score = result.get("score")
        return result.get("returncode") == 0 and isinstance(score, (int, float)) and score > 0
    return False


def check_after_hash(
    errors: list[str], task_id: str, path: Path, expected_hash: object
) -> None:
    label = str(path.relative_to(ROOT)) if path.is_relative_to(ROOT) else str(path)
    if expected_hash in {None, "absent", "ABSENT", "<absent>"}:
        if path.exists():
            errors.append(f"{task_id}: expected absent after write-back: {label}")
        return
    if not isinstance(expected_hash, str) or not SHA256_RE.fullmatch(expected_hash):
        return
    if not path.is_file():
        errors.append(f"{task_id}: after-hash file missing: {label}")
        return
    actual = sha256(path)
    if actual != expected_hash:
        errors.append(
            f"{task_id}: after SHA mismatch for {label}: report={expected_hash} actual={actual}"
        )


def check_before_hash(
    errors: list[str],
    task_id: str,
    path: Path,
    expected_hash: object,
    reported_after: object,
) -> None:
    label = str(path.relative_to(ROOT)) if path.is_relative_to(ROOT) else str(path)
    try:
        repository_relative = path.relative_to(ROOT)
    except ValueError:
        errors.append(f"{task_id}: before-hash path is outside repository: {label}")
        return

    actual_hashes = git_file_sha256_candidates(repository_relative)
    if expected_hash in {None, "absent", "ABSENT", "<absent>"}:
        if actual_hashes:
            errors.append(
                f"{task_id}: expected absent before write-back but HEAD contains {label}"
            )
        return
    if not isinstance(expected_hash, str) or not SHA256_RE.fullmatch(expected_hash):
        errors.append(f"{task_id}: invalid before SHA for {label}: {expected_hash!r}")
        return
    if actual_hashes:
        if expected_hash not in actual_hashes:
            errors.append(
                f"{task_id}: before SHA mismatch for {label}: report={expected_hash} "
                f"HEAD_candidates={sorted(actual_hashes)}"
            )
        return

    # Some legacy task inputs are ignored and have no Git object. They can only
    # be treated as an unchanged pre-existing baseline when before=after and the
    # current file still has that exact digest.
    if reported_after == expected_hash and path.is_file() and sha256(path) == expected_hash:
        return
    errors.append(
        f"{task_id}: before SHA has no verifiable Git baseline for {label}: "
        f"report={expected_hash}"
    )


def audit_hashes(errors: list[str], record: dict, task_root: Path) -> None:
    task_id = record["task_id"]

    def resolve(raw: object) -> Path:
        relative = Path(str(raw))
        if relative.parts and relative.parts[0] == "task":
            return ROOT / relative
        return task_root / relative

    before_after = record.get("sha256_before_after")
    if isinstance(before_after, dict):
        for relative, values in before_after.items():
            if isinstance(values, dict):
                path = resolve(relative)
                check_before_hash(
                    errors, task_id, path, values.get("before"), values.get("after")
                )
                check_after_hash(errors, task_id, path, values.get("after"))

    hashes = record.get("hashes")
    if isinstance(hashes, list):
        for values in hashes:
            if not isinstance(values, dict) or not isinstance(values.get("path"), str):
                continue
            raw_path = Path(values["path"])
            path = (
                ROOT / raw_path
                if raw_path.parts and raw_path.parts[0] == "task"
                else task_root / raw_path
            )
            check_before_hash(
                errors,
                task_id,
                path,
                values.get("before_sha256"),
                values.get("after_sha256"),
            )
            check_after_hash(errors, task_id, path, values.get("after_sha256"))

    repository_hashes = record.get("repository_hashes")
    if isinstance(repository_hashes, dict):
        for relative, values in repository_hashes.items():
            if isinstance(values, dict):
                path = resolve(relative)
                check_before_hash(
                    errors, task_id, path, values.get("before"), values.get("after")
                )
                check_after_hash(errors, task_id, path, values.get("after"))
    elif isinstance(repository_hashes, list):
        for values in repository_hashes:
            if not isinstance(values, dict) or not isinstance(values.get("path"), str):
                continue
            raw_path = Path(values["path"])
            path = raw_path if raw_path.is_absolute() else ROOT / raw_path
            before_hash = values.get("before_sha256", values.get("before"))
            after_hash = values.get("after_sha256", values.get("after"))
            check_before_hash(errors, task_id, path, before_hash, after_hash)
            check_after_hash(
                errors,
                task_id,
                path,
                after_hash,
            )

    after = record.get("sha256_after")
    if isinstance(after, dict):
        for relative, expected_hash in after.items():
            path = ROOT / relative if str(relative).startswith("task/") else task_root / relative
            check_after_hash(errors, task_id, path, expected_hash)

    before = record.get("sha256_before")
    if isinstance(before, dict):
        for relative, expected_hash in before.items():
            path = ROOT / relative if str(relative).startswith("task/") else task_root / relative
            reported_after = after.get(relative) if isinstance(after, dict) else None
            check_before_hash(errors, task_id, path, expected_hash, reported_after)


def claimed_modified_paths(record: dict, task_root: Path) -> set[str]:
    claimed = set()
    for item in record.get("modified_files", []):
        if not isinstance(item, str):
            continue
        cleaned = re.sub(r"\s+\((?:added|removed|modified)\)$", "", item)
        raw_path = Path(cleaned)
        path = ROOT / raw_path if raw_path.parts and raw_path.parts[0] == "task" else task_root / raw_path
        claimed.add(path.relative_to(ROOT).as_posix())
    return claimed


def audit_unblock_snapshot_outputs(errors: list[str]) -> None:
    for report_path in sorted(UNBLOCK_ROOT.glob("*/report.jsonl")):
        snapshot = report_path.parent.name
        try:
            records = [
                json.loads(line)
                for line in report_path.read_text(encoding="utf-8").splitlines()
                if line.strip()
            ]
        except (OSError, UnicodeDecodeError, json.JSONDecodeError) as exc:
            errors.append(f"{snapshot}: cannot parse unblock report for summary audit: {exc}")
            continue

        summary_path = report_path.parent / "summary.json"
        disconnect_path = report_path.parent / "disconnect_verification.json"
        for label, path in (("summary", summary_path), ("disconnect", disconnect_path)):
            if not path.is_file():
                errors.append(f"{snapshot}: missing unblock {label} evidence")

        if summary_path.is_file():
            try:
                summary = json.loads(summary_path.read_text(encoding="utf-8"))
            except (OSError, UnicodeDecodeError, json.JSONDecodeError) as exc:
                errors.append(f"{snapshot}: invalid unblock summary: {exc}")
            else:
                counts = Counter(record.get("status") for record in records)
                if summary.get("total", summary.get("targeted_tasks")) != len(records):
                    errors.append(f"{snapshot}: unblock summary total mismatch")
                for status in ALLOWED_STATUSES:
                    if int(summary.get(status, 0)) != counts[status]:
                        errors.append(f"{snapshot}: unblock summary {status} mismatch")
                modified = [record for record in records if record.get("modified_files")]
                if summary.get("modified_tasks") != len(modified):
                    errors.append(f"{snapshot}: unblock modified-task count mismatch")
                if summary.get("modified_files") != sum(
                    len(record.get("modified_files", [])) for record in modified
                ):
                    errors.append(f"{snapshot}: unblock modified-file count mismatch")

        if disconnect_path.is_file():
            try:
                disconnect = json.loads(disconnect_path.read_text(encoding="utf-8"))
            except (OSError, UnicodeDecodeError, json.JSONDecodeError) as exc:
                errors.append(f"{snapshot}: invalid disconnect evidence: {exc}")
            else:
                if disconnect.get("disconnected") is not True:
                    errors.append(f"{snapshot}: disconnect evidence is not affirmative")


def main() -> int:
    errors = []
    expected = load_expected()
    expected_ids = [task_id for task_id, _ in expected]
    expected_snapshot = dict(expected)
    tasks = discover_tasks()
    records, report_errors = load_reports()
    errors.extend(report_errors)
    audit_unblock_snapshot_outputs(errors)

    record_ids = [record.get("task_id") for record in records]
    counts = Counter(record_ids)
    for task_id in expected_ids:
        if counts[task_id] == 0:
            errors.append(f"{task_id}: missing report record")
        elif counts[task_id] > 1:
            errors.append(f"{task_id}: duplicate report records ({counts[task_id]})")
    for task_id in record_ids:
        if task_id not in expected_snapshot:
            errors.append(f"{task_id}: report record not listed in B组")

    status_counts = Counter()
    snapshot_counts = Counter()
    modified_tasks = 0
    modified_files = 0
    claimed_changes = set()
    for record in records:
        task_id = record.get("task_id")
        if task_id not in expected_snapshot:
            continue
        snapshot = record.get("snapshot")
        status = record.get("status")
        status_counts[status] += 1
        snapshot_counts[(snapshot, status)] += 1

        if snapshot != expected_snapshot[task_id]:
            errors.append(
                f"{task_id}: snapshot mismatch: report={snapshot} workbook={expected_snapshot[task_id]}"
            )
        if status not in ALLOWED_STATUSES:
            errors.append(f"{task_id}: invalid status: {status!r}")
        if not record.get("failure_reason"):
            errors.append(f"{task_id}: missing failure_reason")
        if not record.get("cleanup"):
            errors.append(f"{task_id}: missing cleanup evidence")
        if not any(
            record.get(key)
            for key in ("real_software_generation", "instruction_eval_contract", "instruction")
        ):
            errors.append(f"{task_id}: missing instruction/generation review")

        modified = record.get("modified_files")
        if not isinstance(modified, list):
            errors.append(f"{task_id}: modified_files is not a list")
        elif modified:
            modified_tasks += 1
            modified_files += len(modified)
        if status != "passed" and modified:
            evaluator_only = status == "environment_blocked" and all(
                isinstance(path, str) and Path(path).name == "eval.py"
                for path in modified
            )
            if not evaluator_only:
                errors.append(
                    f"{task_id}: blocked record claims non-evaluator repository modifications"
                )
        if status == "passed" and not eval_pass_evidence(record):
            errors.append(f"{task_id}: passed without recognizable successful formal eval evidence")

        task_info = tasks.get(task_id)
        if task_info is None:
            errors.append(f"{task_id}: task definition not found")
        else:
            task_root, task_data = task_info
            if task_data.get("snapshot") != snapshot:
                errors.append(f"{task_id}: task JSON snapshot differs from report")
            audit_hashes(errors, record, task_root)
            claimed_changes.update(claimed_modified_paths(record, task_root))

        for raw_path in sorted(set(path_strings(record))):
            path = ROOT / raw_path.rstrip("/")
            if not path.exists():
                errors.append(f"{task_id}: referenced log path missing: {raw_path}")

    actual_changes = git_task_changes()
    for path in sorted(actual_changes - claimed_changes):
        errors.append(f"unclaimed Task worktree change: {path}")
    for path in sorted(claimed_changes - actual_changes):
        errors.append(f"reported modified file is clean or missing from Git status: {path}")

    summary = {
        "expected": len(expected_ids),
        "reported": len(records),
        "status_counts": dict(sorted(status_counts.items(), key=lambda item: str(item[0]))),
        "modified_tasks": modified_tasks,
        "modified_files_reported": modified_files,
        "by_snapshot": {
            snapshot: {
                status: snapshot_counts[(snapshot, status)]
                for status in sorted(ALLOWED_STATUSES)
                if snapshot_counts[(snapshot, status)]
            }
            for snapshot in sorted({snapshot for _, snapshot in expected})
        },
        "errors": len(errors),
    }
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    if errors:
        print("\nAUDIT ERRORS", file=sys.stderr)
        for error in errors:
            print(f"- {error}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
